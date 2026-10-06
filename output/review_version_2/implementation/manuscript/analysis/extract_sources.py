"""Read-only provenance extraction; preserves workbook values and hyperlink targets."""
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from urllib.parse import unquote, urlsplit

import openpyxl

ROOT = Path(__file__).parent
WORKBOOK = ROOT.parent / 'supplementary/Perspectives.xlsx'
GOLD = ROOT.parent / 'supplementary/gold_standard.json'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump_json(name, value):
    (ROOT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def dump_csv(name, rows):
    with (ROOT / name).open('w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        for row in rows:
            writer.writerow({k: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v for k, v in row.items()})


def urls(value):
    return re.findall(r'https?://[^\s<>]+', value or '')


def main():
    before = sha(WORKBOOK)
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=False)
    gold = json.loads(GOLD.read_text())
    inventory = []
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.value is None and cell.hyperlink is None and cell.comment is None:
                    continue
                inventory.append({
                    'sheet': sheet.title, 'cell': cell.coordinate, 'value': cell.value,
                    'data_type': cell.data_type,
                    'hyperlink_target': cell.hyperlink.target if cell.hyperlink else None,
                    'hyperlink_location': cell.hyperlink.location if cell.hyperlink else None,
                    'comment_text': cell.comment.text if cell.comment else None,
                    'fill_type': cell.fill.fill_type, 'fill_color_type': cell.fill.fgColor.type,
                    'fill_color': cell.fill.fgColor.index,
                })
    records, links = [], []
    for topic in gold['topics']:
        sheet = workbook[{'weird': '01_WEIRD', 'serbian': '02_Serbian'}[topic['arm']]]
        row = topic['pair_id'] + 1
        for ordinal, perspective in enumerate(topic['perspectives'], 1):
            col = 2 + 3 * (ordinal - 1)
            explanation, summary, source = [sheet.cell(row, col + offset) for offset in range(3)]
            value_urls = urls(source.value)
            hyperlink = source.hyperlink.target if source.hyperlink else None
            source_urls = list(dict.fromkeys(value_urls + ([hyperlink] if hyperlink else [])))
            flags = []
            if hyperlink and hyperlink not in value_urls:
                if any(url.split('#', 1)[0] == hyperlink.split('#', 1)[0] for url in value_urls):
                    flags.append('displayed_url_and_hyperlink_differ_only_by_fragment')
                else:
                    flags.append('displayed_url_and_hyperlink_target_conflict')
            if not source_urls:
                flags.append('textual_source_without_url')
            if source.fill.fill_type == 'solid' and source.fill.fgColor.index == 'FFFF0000':
                flags.append('red_fill_in_original_workbook_meaning_unspecified')
            embedded_url = None
            for url in source_urls:
                if urlsplit(url).netloc == 'claude.ai':
                    flags.append('private_chat_url_in_source_cell')
                    if '#:~:text=' in url:
                        decoded = unquote(url.split('#:~:text=', 1)[1])
                        targets = re.findall(r'https?://[^\s<>]+', decoded)
                        if targets:
                            embedded_url = targets[0]
                links.append({'perspective_id': perspective['perspective_id'], 'sheet': sheet.title,
                              'source_cell': source.coordinate, 'url': url,
                              'in_displayed_value': url in value_urls, 'is_hyperlink_target': url == hyperlink,
                              'support_verification': 'not_verified_by_extraction'})
            item = {
                'perspective_id': perspective['perspective_id'], 'topic_id': topic['topic_id'],
                'pair_id': topic['pair_id'], 'arm': topic['arm'], 'ordinal': ordinal,
                'gold_topic_title': topic['topic_title'], 'workbook_topic_title': sheet.cell(row, 1).value,
                'sheet': sheet.title, 'topic_cell': f'A{row}',
                'explanation_cell': explanation.coordinate, 'summary_cell': summary.coordinate,
                'source_cell': source.coordinate, 'workbook_explanation': explanation.value,
                'workbook_summary': summary.value, 'source_text_exact': source.value,
                'hyperlink_target_exact': hyperlink, 'urls_from_value': value_urls,
                'all_recorded_urls': source_urls, 'source_comment': source.comment.text if source.comment else None,
                'source_fill_type': source.fill.fill_type, 'source_fill_color': source.fill.fgColor.index,
                'summary_exact_match': perspective['summary'] == summary.value,
                'explanation_exact_match': perspective['explanation'] == explanation.value,
                'matching_confidence': 'exact_summary_and_explanation',
                'extremity_tier_from_gold_only': perspective.get('extremity_tier'),
                'workbook_extremity_tier': None, 'flags': flags,
                'public_url_decoded_from_chat_fragment_unverified': embedded_url,
                'support_verification': 'source_record_present_not_full_support_validation',
            }
            assert item['summary_exact_match'] and item['explanation_exact_match'], item['perspective_id']
            records.append(item)
    assert len(records) == 80 and len({r['perspective_id'] for r in records}) == 80
    assert all(r['source_text_exact'] for r in records)
    assert sha(WORKBOOK) == before, 'Read-only extraction changed the workbook'
    summary = {
        'workbook_path': str(WORKBOOK), 'workbook_sha256': before, 'gold_path': str(GOLD),
        'gold_sha256': sha(GOLD), 'sheets': workbook.sheetnames, 'populated_or_annotated_cells': len(inventory),
        'reference_count': len(records), 'topic_count': len(gold['topics']),
        'exact_summary_matches': sum(r['summary_exact_match'] for r in records),
        'exact_explanation_matches': sum(r['explanation_exact_match'] for r in records),
        'populated_source_cells': sum(bool(r['source_text_exact']) for r in records),
        'source_cells_with_at_least_one_recorded_url': sum(bool(r['all_recorded_urls']) for r in records),
        'source_cells_without_url': sum(not r['all_recorded_urls'] for r in records),
        'source_cells_with_hyperlink_object': sum(bool(r['hyperlink_target_exact']) for r in records),
        'reference_url_associations': len(links), 'distinct_recorded_urls': len({r['url'] for r in links}),
        'url_conflict_count': sum('displayed_url_and_hyperlink_target_conflict' in r['flags'] for r in records),
        'fragment_only_url_difference_count': sum('displayed_url_and_hyperlink_differ_only_by_fragment' in r['flags'] for r in records),
        'private_chat_url_count': sum('private_chat_url_in_source_cell' in r['flags'] for r in records),
        'comment_count': sum(r['comment_text'] is not None for r in inventory),
        'tiers_present_in_workbook': False, 'gold_tier_counts': dict(Counter(r['extremity_tier_from_gold_only'] for r in records)),
        'source_support_validation': 'No blanket validation. See SOURCE_AUDIT.md for bounded web checks.',
    }
    dump_json('workbook_cells_exact.json', inventory)
    dump_json('source_provenance_ledger.json', records)
    dump_csv('source_provenance_ledger.csv', records)
    dump_csv('source_urls.csv', links)
    dump_json('extraction_summary.json', summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
